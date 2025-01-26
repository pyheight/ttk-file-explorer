import os
import time 
import threading 
import ctypes
from ctypes import wintypes

import zipfile 
import winshell
from win32com.shell import shell
from PIL import Image   
from docx import Document    
from openpyxl import Workbook    
from fpdf import FPDF    
from pptx import Presentation   

from pathlib import Path
import pythoncom

from model.system_info import get_system_encoding


def perform_file_operation(files, target, mode):
    pythoncom.CoInitialize()
    fo = pythoncom.CoCreateInstance(shell.CLSID_FileOperation, None, 
                                    pythoncom.CLSCTX_INPROC_SERVER,
                                    shell.IID_IFileOperation)

    for file in files:
        src = shell.SHCreateItemFromParsingName(file, None, shell.IID_IShellItem)
        if mode != 'recycle':
            dst = shell.SHCreateItemFromParsingName(target, None, shell.IID_IShellItem)
        if mode == 'copy':
            fo.CopyItem(src, dst, None, None)
        elif mode == 'move':
            fo.MoveItem(src, dst, None, None)
        elif mode == 'recycle':
            fo.DeleteItem(src, None)  

    fo.PerformOperations()
    pythoncom.CoUninitialize() 


def copy_file(source, target):
    threading.Thread(target=perform_file_operation, args=(source, target, 'copy',)).start()


def move_file(source, target):
    threading.Thread(target=perform_file_operation, args=(source, target, 'move',)).start()


def delete_file(source, hwnd):
    winshell.delete_file(source, allow_undo=False, hWnd=hwnd)


def recycle_file(source):
    threading.Thread(target=perform_file_operation, args=(source, None, 'recycle',)).start()


def rename_file(source, target):
    threading.Thread(target=winshell.rename_file, args=(source, target,)).start()


def new_shortcut(target, src, arguments='', description=''):
    winshell.CreateShortcut(
        Path=target,
        Target=src,
        Arguments=arguments,
        Icon=(src, 0),
        Description=description,
    )


def get_shortcut_path(path):
    return winshell.shortcut(path).path  


def search_file(search_path, pattern):
    search_path = Path(search_path)  
    matching_files = list(search_path.rglob(pattern))  
    matching_folders = [str(p) for p in search_path.rglob(f'*/{pattern}') if p.is_dir()]  
    
    return matching_files + matching_folders  


def create_file(path, file_type):  
    if file_type == 'folder':  
        os.makedirs(path)
    elif file_type in ('.bmp', '.png'):  
        image = Image.new('RGB', (1, 1))    
        image.save(path, format=file_type[1:].upper()) 
    elif file_type in ('.docx', '.doc'):  
        Document().save(path)    
    elif file_type == '.pdf':   
        pdf = FPDF()    
        pdf.add_page()   
        pdf.output(path, 'F')   
    elif file_type in ('.pptx', '.ppt'):    
        Presentation().save(path)
    elif file_type in ('.xlsx', '.xls'):    
        Workbook().save(path)    
    elif file_type == '.zip':    
        with zipfile.ZipFile(path, 'w'):  
            pass   
    else: 
        with open(path, 'w', encoding='utf-8') as f:    
            f.write('')


SHGetPathFromIDList = ctypes.windll.shell32.SHGetPathFromIDListW
SHGetPathFromIDList.argtypes = [ctypes.c_void_p, wintypes.LPWSTR]
SHGetPathFromIDList.restype = wintypes.BOOL
SHBrowseForFolder = ctypes.windll.shell32.SHBrowseForFolderW
SHBrowseForFolder.argtypes = [ctypes.c_void_p]
SHBrowseForFolder.restype = ctypes.c_void_p
CoTaskMemFree = ctypes.windll.ole32.CoTaskMemFree
CoTaskMemFree.argtypes = [ctypes.c_void_p]
CoTaskMemFree.restype = None
MAX_PATH = wintypes.MAX_PATH  
BROWSEINFO_FLAGS = 0x0001 | 0x0020 | 0x0010 | 0x8000 | 0x4000 | 0x0040 | 0x0080 | 0x0100 

SEE_MASK = 0x00000040 | 0x0000000C
ShellExecuteEx = ctypes.windll.shell32.ShellExecuteEx  
ShellExecuteEx.restype = wintypes.BOOL  
system_encod = get_system_encoding()


class BROWSEINFO(ctypes.Structure):
    _fields_ = [
        ('hwndOwner', wintypes.HWND),
        ('pidlRoot', ctypes.c_void_p),
        ('pszDisplayName', ctypes.c_wchar_p),
        ('lpszTitle', wintypes.LPCWSTR),
        ('ulFlags', wintypes.UINT),
        ('lpfn', ctypes.c_void_p),
        ('lParam', ctypes.c_ulong),
        ('iImage', wintypes.INT),
    ]


def tree_askdirectory(window, title):  
    bInfo = BROWSEINFO()
    bInfo.hwndOwner = window
    bInfo.lpszTitle = title
    bInfo.ulFlags = BROWSEINFO_FLAGS

    pidl = SHBrowseForFolder(ctypes.byref(bInfo))
    if pidl:
        try:
            path = ctypes.create_unicode_buffer(wintypes.MAX_PATH)
            SHGetPathFromIDList(pidl, path)
            selected_folder = path.value
        finally:
            CoTaskMemFree(pidl)
    else:
        selected_folder = None

    return selected_folder


class SHELLEXECUTEINFO(ctypes.Structure):  
    _fields_ = (  
        ('cbSize', wintypes.DWORD),  
        ('fMask', ctypes.c_ulong),  
        ('hwnd', wintypes.HANDLE),  
        ('lpVerb', ctypes.c_char_p),  
        ('lpFile', ctypes.c_char_p),  
        ('lpParameters', ctypes.c_char_p),  
        ('lpDirectory', ctypes.c_char_p),  
        ('nShow', ctypes.c_int),  
        ('hInstApp', wintypes.HINSTANCE),  
        ('lpIDList', ctypes.c_void_p),  
        ('lpClass', ctypes.c_char_p),  
        ('hKeyClass', wintypes.HKEY),  
        ('dwHotKey', wintypes.DWORD),  
        ('hIconOrMonitor', wintypes.HANDLE),  
        ('hProcess', wintypes.HANDLE),  
    )  


def open_with(path, mode, sleep_time=86400):  
    def shell_execute():
        sei = SHELLEXECUTEINFO()  
        sei.cbSize = ctypes.sizeof(sei)  
        sei.fMask = SEE_MASK
        sei.lpVerb = mode.encode(system_encod)  
        sei.lpFile = path.encode(system_encod)
        sei.nShow = 1  
        try:
            if ShellExecuteEx(ctypes.byref(sei)):  
                time.sleep(sleep_time)
        except Exception:
            return
    threading.Thread(target=shell_execute, daemon=True).start()
